"""The completed workbook has to be readable, not just correct.

The supplied template sets a width for three columns only. Everything from
Candidate Evidence rightwards falls back to Excel's ~8-character default while
carrying 170+ characters of wrapped text, which renders as a thin ribbon of
clipped content. These tests keep the formatting pass in place.
"""
from __future__ import annotations

import pathlib
import sys

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))

from asset_factory.qa import COLUMN_WIDTHS, WRAP_COLUMNS, fill  # noqa: E402

ROOT = pathlib.Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "inputs" / "qa-checklist.xlsx"


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    if not TEMPLATE.exists():
        pytest.skip("QA checklist template not present")
    target = tmp_path_factory.mktemp("qa") / "completed.xlsx"
    fill(TEMPLATE, target)
    wb = openpyxl.load_workbook(target)
    return wb[wb.sheetnames[0]]


def _columns(sheet) -> dict[str, int]:
    for row in range(1, 11):
        values = {
            str(sheet.cell(row, col).value or "").strip().lower(): col
            for col in range(1, sheet.max_column + 1)
        }
        if "qa check" in values:
            return {k: v for k, v in values.items() if k}
    raise AssertionError("header row not found")


def test_every_column_has_an_explicit_width(sheet):
    for name, index in _columns(sheet).items():
        letter = get_column_letter(index)
        dim = sheet.column_dimensions.get(letter)
        assert dim and dim.width, f"column {letter} ({name}) has no width set"


def test_evidence_column_is_wide_enough_to_read(sheet):
    columns = _columns(sheet)
    letter = get_column_letter(columns["candidate evidence"])
    assert sheet.column_dimensions[letter].width >= 60


def test_long_text_columns_wrap(sheet):
    columns = _columns(sheet)
    for name in WRAP_COLUMNS:
        if name not in columns:
            continue
        cell = sheet.cell(3, columns[name])
        assert cell.alignment.wrap_text, f"{name} does not wrap"
        assert cell.alignment.vertical == "top", f"{name} is not top-aligned"


def test_rows_are_not_height_locked(sheet):
    """A fixed height clips wrapped text. Excel must be free to auto-fit."""
    locked = [r for r, d in sheet.row_dimensions.items()
              if d.height and r >= 3]
    assert not locked, f"rows with fixed heights would clip content: {locked}"


def test_header_is_frozen(sheet):
    assert sheet.freeze_panes, "header row is not frozen"


def test_width_map_covers_the_supplied_template(sheet):
    for name in _columns(sheet):
        assert name in COLUMN_WIDTHS, f"no width defined for column {name!r}"
