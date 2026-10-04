"""Reads the Site Variable Sheet and produces a validated variable set."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .model import (
    Reason,
    Status,
    Variable,
    advisory_validator_for,
    blocking_validator_for,
    normalize_status,
)

SITE_SHEET_CANDIDATES = ("East County Variables",)
HEADER_TOKENS = ("variable", "category")

# "Value A: Request a Quote\nValue B: Request a Pest Inspection\nExpected ..."
_CANDIDATE_RE = re.compile(r"^\s*Value\s+([A-Z])\s*:\s*(.+?)\s*$", re.M)
_EXPECTED_RE = re.compile(r"^\s*Expected candidate action\s*:\s*(.+?)\s*$", re.M | re.I)


@dataclass
class VariableSet:
    variables: dict[str, Variable]
    source: str

    def __getitem__(self, key: str) -> Variable:
        return self.variables[key]

    def get(self, key: str) -> Variable | None:
        return self.variables.get(key)

    def approved(self, key: str) -> str | None:
        """The value to write, or None if this field must not be written."""
        var = self.variables.get(key)
        return var.value if var and var.applicable else None

    @property
    def applicable(self) -> list[Variable]:
        return [v for v in self.variables.values() if v.applicable]

    @property
    def exceptions(self) -> list[Variable]:
        return [v for v in self.variables.values() if not v.applicable]

    def summary(self) -> dict[str, int]:
        counts: dict[str, int] = {"total": len(self.variables), "applicable": 0}
        for var in self.variables.values():
            if var.applicable:
                counts["applicable"] += 1
            else:
                reason = var.primary_reason.value if var.primary_reason else "unknown"
                counts[reason] = counts.get(reason, 0) + 1
        return counts


def _find_header_row(ws) -> int:
    for row in range(1, min(ws.max_row, 10) + 1):
        cells = [str(ws.cell(row, c).value or "").strip().lower() for c in range(1, 5)]
        if any(t in cells[0] for t in HEADER_TOKENS) and "category" in cells[1]:
            return row
    raise ValueError("could not locate the header row in the variable sheet")


def _split_conflict(text: str) -> tuple[list[str], str | None]:
    candidates = [m.group(2).strip() for m in _CANDIDATE_RE.finditer(text)]
    expected = _EXPECTED_RE.search(text)
    return candidates, expected.group(1).strip() if expected else None


def load_variable_sheet(path: str | Path, sheet_name: str | None = None) -> VariableSet:
    wb = openpyxl.load_workbook(path, data_only=True)

    if sheet_name is None:
        for name in SITE_SHEET_CANDIDATES:
            if name in wb.sheetnames:
                sheet_name = name
                break
        else:  # fall back to the last sheet, which holds site-specific values
            sheet_name = wb.sheetnames[-1]

    ws = wb[sheet_name]
    header = _find_header_row(ws)
    variables: dict[str, Variable] = {}

    for row in range(header + 1, ws.max_row + 1):
        key = ws.cell(row, 1).value
        if key is None or not str(key).strip():
            continue
        key = str(key).strip()

        var = Variable(
            key=key,
            category=str(ws.cell(row, 2).value or "").strip(),
            raw_value=ws.cell(row, 3).value,
            status=normalize_status(ws.cell(row, 4).value),
            row=row,
        )
        _classify(var)
        variables[key] = var

    return VariableSet(variables=variables, source=str(path))


def _classify(var: Variable) -> None:
    """Apply the demotion rules. This function never makes a field usable."""

    # 1. Non-approved statuses can never be written, whatever the cell holds.
    if not var.status.is_usable:
        if var.status is Status.CONFLICTING and var.has_text:
            candidates, expected = _split_conflict(var.value or "")
            var.candidates = candidates
            detail = (
                f"{len(candidates)} candidate values supplied; none selected"
                if candidates
                else "multiple values supplied in a single cell"
            )
            if expected:
                detail += f" -- sheet instructs: {expected}"
            var.note(Reason.MULTIPLE_CANDIDATES, detail)
        elif var.has_text:
            # A value is sitting in the cell but carries no approval. This is
            # the dangerous case: naive automation keyed on truthiness writes it.
            var.note(
                Reason.PHANTOM_VALUE,
                f"cell contains {var.value!r} but status is {var.status.value!r}; "
                "value is present without approval and must not be written",
            )
        else:
            var.note(
                {
                    Status.MISSING: Reason.NOT_SUPPLIED,
                    Status.NEEDS_CONFIRMATION: Reason.AWAITING_CONFIRMATION,
                }.get(var.status, Reason.UNKNOWN_STATUS),
                f"no value supplied; status is {var.status.value!r}",
            )
        return

    # 2. Approved but empty -- treat as missing, do not write an empty string.
    if not var.has_text:
        var.demote(
            Reason.NOT_SUPPLIED,
            "status is 'Approved' but the cell is empty; nothing to write",
        )
        return

    # 3. Approved and populated. A blocking validator can still demote it;
    #    an advisory one records a concern but leaves the value writable.
    blocking = blocking_validator_for(var.key)
    if blocking:
        problem = blocking(var.value or "")
        if problem:
            var.demote(Reason.FAILED_FORMAT, problem)
            return

    advisory = advisory_validator_for(var.key)
    if advisory:
        concern = advisory(var.value or "")
        if concern:
            var.advise(concern)


def to_json(vs: VariableSet) -> str:
    payload = {
        "source": vs.source,
        "summary": vs.summary(),
        "variables": [
            {
                "key": v.key,
                "category": v.category,
                "status": v.status.value,
                "value": v.value if v.applicable else None,
                "raw_value": None if v.raw_value is None else str(v.raw_value),
                "applicable": v.applicable,
                "candidates": v.candidates,
                "findings": [
                    {
                        "reason": f.reason.value,
                        "detail": f.detail,
                        "severity": f.severity.value,
                    }
                    for f in v.findings
                ],
                "row": v.row,
            }
            for v in vs.variables.values()
        ],
    }
    return json.dumps(payload, indent=2)
